"""Build a comprehensive English study guide for the presentation.

Run:
    python build_study_guide.py

Output:
    Cross-Modal_Presentation_Study_Guide.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ── Colors ───────────────────────────────────────────────────────────────────
INDIGO  = RGBColor(0x63, 0x66, 0xf1)
VIOLET  = RGBColor(0x8b, 0x5c, 0xf6)
TEXT    = RGBColor(0x1f, 0x29, 0x37)
MUTED   = RGBColor(0x6b, 0x71, 0x85)
ACCENT  = RGBColor(0x10, 0xb9, 0x81)
DANGER  = RGBColor(0xef, 0x44, 0x44)
LIGHT   = RGBColor(0xf3, 0xf4, 0xf6)

doc = Document()

# ── Page margins ─────────────────────────────────────────────────────────────
for s in doc.sections:
    s.top_margin    = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin   = Cm(2.2)
    s.right_margin  = Cm(2.2)

# ── Default style ────────────────────────────────────────────────────────────
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = TEXT


# ── Helpers ──────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = INDIGO
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = VIOLET
    return p


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = TEXT
    return p


def para(text, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p


def rich(parts):
    """parts = [(text, {bold, italic, color, mono})]"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for text, attrs in parts:
        r = p.add_run(text)
        r.font.size   = Pt(11)
        r.font.bold   = attrs.get("bold", False)
        r.font.italic = attrs.get("italic", False)
        if attrs.get("mono"):
            r.font.name = "Consolas"
        if attrs.get("color") is not None:
            r.font.color.rgb = attrs["color"]
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + 0.6 * level)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def code_block(code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.4)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(10)
    r = p.add_run(code)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    # Light gray background via paragraph shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F3F4F6")
    pPr.append(shd)
    return p


def callout(title, body, color_hex="6366F1"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_bg(cell, "EEF2FF")
    # Border via colored left edge — approximated with bold title
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    r = p.add_run(f"💡 {title}")
    r.font.bold = True
    r.font.color.rgb = INDIGO
    r.font.size = Pt(11)
    p2 = cell.add_paragraph(body)
    p2.runs[0].font.size = Pt(10.5)
    doc.add_paragraph()


def make_table(header, rows, header_bg="6366F1"):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Light Grid Accent 1"
    # Header row
    for j, txt in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(txt)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        r.font.size = Pt(11)
        set_cell_bg(c, header_bg)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10.5)
    doc.add_paragraph()
    return t


def page_break():
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
# COVER PAGE
# ════════════════════════════════════════════════════════════════════════════
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(120)
r = cover.add_run("Cross-Modal Generation\n")
r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = INDIGO
r = cover.add_run("Presentation Study Guide\n")
r.font.size = Pt(22); r.font.color.rgb = VIOLET

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_before = Pt(20)
r = sub.add_run("Image Captioning + CLIP Retrieval\n")
r.font.size = Pt(14); r.font.italic = True; r.font.color.rgb = MUTED
r = sub.add_run("ResNet18 · Transformer Decoder · CLIP ViT-B/32 · Flickr8k")
r.font.size = Pt(11); r.font.color.rgb = MUTED

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(80)
r = meta.add_run("7-minute presentation · 2-day prep guide\n")
r.font.size = Pt(11); r.font.color.rgb = MUTED
r = meta.add_run("Generated 2026-05-16")
r.font.size = Pt(10); r.font.color.rgb = MUTED

page_break()

# ════════════════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS
# ════════════════════════════════════════════════════════════════════════════
h1("Table of Contents")
toc = [
    "1.  How to Use This Guide",
    "2.  Your 2-Day Study Plan",
    "3.  The 7-Slide Template (At a Glance)",
    "4.  Project Overview (What You Built)",
    "5.  Glossary of Key Terms",
    "6.  Architecture Deep Dive",
    "      6.1  The Captioning Model — ResNet18 + Transformer",
    "      6.2  CLIP for Retrieval — Why It Works Without Training",
    "7.  Training Details & Results",
    "8.  Loss Curve — How to Read & Explain It",
    "9.  Slide-by-Slide Speaker Notes",
    "10. Anticipated Q&A",
    "11. Live Demo Checklist",
    "12. The Night Before — Quick Review",
]
for entry in toc:
    bullet(entry)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 1. HOW TO USE
# ════════════════════════════════════════════════════════════════════════════
h1("1. How to Use This Guide")

para(
    "This guide turns the project you built into a presentation you can confidently deliver in 7 minutes. "
    "It assumes the audience is a mixed group — some technical, some not — and that you have already returned "
    "to coding recently. Every technical term is defined the first time it appears."
)

h3("Read in this order:")
bullet("Day 1 morning  → Sections 3, 4, 5  (the big picture + vocabulary)")
bullet("Day 1 evening  → Section 6  (architecture deep dive — the heart of the talk)")
bullet("Day 2 morning  → Sections 7, 8, 9  (training, loss curve, speaker notes)")
bullet("Day 2 evening  → Section 10  (Q&A practice) + run the demo end-to-end")
bullet("Right before talk → Section 12  (10-minute final review)")

callout(
    "One golden rule",
    "If you only remember one sentence: \"We built two halves — captioning is a model we trained ourselves "
    "from scratch (ResNet18 + Transformer); retrieval uses CLIP, a model OpenAI pre-trained on 400 million "
    "image-text pairs.\" Everything else hangs off this."
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 2. STUDY PLAN
# ════════════════════════════════════════════════════════════════════════════
h1("2. Your 2-Day Study Plan")

para("A realistic schedule for someone returning to ML. Adjust the times to your day.")

h2("Day 1 — Understand")

make_table(
    ["Time", "Task", "Goal"],
    [
        ["45 min", "Read Sections 3–5",                    "Internalize the high-level story + vocabulary"],
        ["60 min", "Read Section 6 (architecture)",        "Be able to draw the diagram from memory"],
        ["30 min", "Open demo.py, trace the code",         "Connect each module to the concepts you just read"],
        ["20 min", "Run the demo locally",                 "Make sure it works on your machine before showing it"],
    ],
)

h2("Day 2 — Rehearse")

make_table(
    ["Time", "Task", "Goal"],
    [
        ["45 min", "Read Sections 7–9, study loss curve",  "Understand metrics + own the speaker notes"],
        ["30 min", "Write/refine your slides",             "Use the template in Section 3"],
        ["45 min", "Deliver out loud 3× with a timer",     "Hit the 7-minute mark consistently"],
        ["30 min", "Practice Q&A (Section 10) with a friend", "Comfortable improvising on hard questions"],
        ["15 min", "Set up demo on the presentation laptop", "Avoid live disasters"],
    ],
)

callout(
    "Tip",
    "Record your first run-through on your phone and play it back. You will catch filler words ('um', 'so') and "
    "places where you talk too fast. This single trick is worth more than 10 extra read-throughs."
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 3. THE 7-SLIDE TEMPLATE
# ════════════════════════════════════════════════════════════════════════════
h1("3. The 7-Slide Template — At a Glance")

para("This is the entire structure. Each slide has one main idea. Detailed speaker notes are in Section 9.")

make_table(
    ["#", "Slide", "Time", "Main Idea"],
    [
        ["1", "Title & Problem",          "~30 s",  "What we built, in one sentence"],
        ["2", "Motivation",               "~45 s",  "Why bridging image ↔ text matters"],
        ["3", "Architecture Overview",    "~60 s",  "Two halves: captioning + retrieval"],
        ["4", "The Captioning Model",     "~75 s",  "ResNet18 encoder + Transformer decoder + cross-attention"],
        ["5", "Training & Results",       "~60 s",  "Flickr8k, AdamW, best val loss 2.98 at epoch 12"],
        ["6", "CLIP for Retrieval",       "~60 s",  "Pretrained joint embedding + cosine similarity"],
        ["7", "Demo + Closing",           "~90 s",  "Live demo, summary, future work"],
    ],
)

callout(
    "Slide design rules",
    "(1) One idea per slide.  (2) No more than 6 bullet points.  (3) Font ≥ 24 pt.  "
    "(4) Diagrams beat text every time.  (5) Save your most important sentence for the closing slide."
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 4. PROJECT OVERVIEW
# ════════════════════════════════════════════════════════════════════════════
h1("4. Project Overview — What You Built")

h2("The 30-second pitch")
para(
    "A web application that connects vision and language in two directions. Drag an image in → it generates a "
    "natural-language caption. Type a description → it retrieves the three most similar images from the dataset. "
    "Built with PyTorch and Flask; the entire frontend is plain HTML, CSS, and JavaScript."
)

h2("System diagram")
code_block(
"""IMAGE  ──┐                                              ┌──→ "a dog runs on grass"
         ├──► [ResNet18] ──► 49 image tokens ──► [Transformer Decoder] ──┘
         │                                       (autoregressive)
         │                                                              CAPTIONING
─────────┼─────────────────────────────────────────────────────────────────────────
TEXT  ───┤                                              ┌──→ top-3 images + captions
         ├──► [CLIP Text Enc.]  ───┐                    │
         │                          ├──► cosine sim ────┘
         └──► [CLIP Image Enc.] ───┘                              RETRIEVAL""")

h2("Three pieces of code")
make_table(
    ["File", "Role"],
    [
        ["train.ipynb",         "Notebook that loads Flickr8k, defines the model, runs 20 training epochs, "
                                "saves the best checkpoint."],
        ["demo.py",             "Flask backend. Loads the trained captioning checkpoint + pretrained CLIP, "
                                "exposes three routes (GET /, POST /caption, POST /retrieve)."],
        ["templates/index.html","Single-page frontend with two sections — text retrieval (3 result cards) and "
                                "image upload (drag-and-drop) — using vanilla HTML/CSS/JS."],
    ],
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 5. GLOSSARY
# ════════════════════════════════════════════════════════════════════════════
h1("5. Glossary of Key Terms")

para(
    "Every term you may need on stage, with a short definition and an analogy. "
    "If a term feels fuzzy, come back here.",
    italic=True, color=MUTED,
)

glossary = [
    ("Modality",
     "A type of data — image, text, audio, video are all modalities.",
     "Different languages."),
    ("Cross-modal",
     "Anything that bridges two modalities — e.g. image → text (captioning) or text → image (retrieval).",
     "Like translating between languages."),
    ("Embedding",
     "A vector (list of numbers) that represents an object — a sentence, an image — in a way the model can compare.",
     "A coordinate in a high-dimensional space."),
    ("Encoder",
     "A neural network that turns an input (image or text) into an embedding.",
     "A 'meaning extractor.'"),
    ("Decoder",
     "A neural network that turns an embedding back into structured output (e.g. a sentence).",
     "A 'narrator.'"),
    ("Token",
     "The smallest unit the model sees — a word or sub-word (e.g. 'play' + '##ing').",
     "A puzzle piece."),
    ("Tokenizer",
     "The component that splits raw text into tokens (we use BERT's WordPiece tokenizer).",
     "A scissors that cuts text into model-friendly chunks."),
    ("Attention",
     "A mechanism that lets the model decide which inputs to focus on when producing each output.",
     "A spotlight that moves around the scene."),
    ("Self-attention",
     "Tokens in a sequence attend to each other — the decoder uses this to consider all prior generated words.",
     "Each word looking back at the words before it."),
    ("Cross-attention",
     "Tokens in one sequence attend to a DIFFERENT sequence — here, decoder words attend to image tokens.",
     "The decoder peeking at the image while writing."),
    ("Causal mask",
     "An upper-triangular mask that prevents a token from attending to future positions during training.",
     "Covering future words so the model can't cheat."),
    ("Autoregressive",
     "Generating one token at a time, each conditioned on all previous tokens.",
     "Writing word by word, never looking ahead."),
    ("Teacher forcing",
     "During training, the decoder is fed the GROUND-TRUTH previous word, not its own prediction. "
     "This makes training parallelizable and stable.",
     "Showing the student the correct answer at every step so they learn faster."),
    ("Greedy decoding",
     "At inference, pick the single most likely next token at every step.",
     "Always go for the safest bet — fast but not always best."),
    ("Cross-entropy loss",
     "The standard loss for classification: how far off was the predicted probability distribution from the truth?",
     "How surprised was the model by the right answer?"),
    ("Perplexity (PPL)",
     "exp(loss). Intuitive measure: 'on average, how many words is the model torn between?' Lower is better.",
     "If PPL = 20, the model is choosing the correct word from ~20 plausible candidates."),
    ("Epoch",
     "One full pass through the training dataset.",
     "Reading a textbook cover-to-cover once."),
    ("AdamW",
     "Adam optimizer with decoupled weight decay. The default modern optimizer for transformers.",
     "A smart auto-pilot for gradient descent."),
    ("Cosine similarity",
     "The cosine of the angle between two vectors. 1 = same direction, 0 = perpendicular, -1 = opposite.",
     "How aligned are these two arrows?"),
    ("Weight tying",
     "Reusing the input embedding matrix as the output projection. Halves the head's parameter count and "
     "often improves generalization.",
     "One employee doing two related jobs."),
    ("ResNet18",
     "An 18-layer convolutional neural network from 2015 with residual (skip) connections. Pretrained on "
     "ImageNet here.",
     "A solid, mid-sized image encoder — the 'Honda Civic' of CNNs."),
    ("CLIP",
     "Contrastive Language-Image Pre-training. OpenAI 2021. Trained on 400 M (image, caption) pairs to put "
     "images and texts in a SHARED embedding space.",
     "A model that learned to match photos to captions, and now we reuse that skill."),
    ("Joint embedding space",
     "A vector space where embeddings from two modalities (image + text) live side by side and can be compared.",
     "One coordinate system where photos and sentences are dots that cluster by meaning."),
    ("Overfitting",
     "When the model memorizes training data instead of learning to generalize — val loss starts rising "
     "while train loss keeps dropping.",
     "A student who memorizes past exam answers but fails the new exam."),
    ("Early stopping",
     "Saving the model from the epoch with the best validation loss and discarding later, overfit weights.",
     "Knowing when to stop tweaking."),
    ("Flickr8k",
     "A small captioning dataset: 8,091 images, 5 human-written captions each.",
     "The 'starter' captioning benchmark — small but clean."),
    ("Flask",
     "A lightweight Python web framework. We use it to expose three HTTP routes for the demo.",
     "A minimal toolkit for turning Python functions into web endpoints."),
]

for term, defn, analogy in glossary:
    rich([
        (f"{term}  ", {"bold": True, "color": INDIGO}),
        (f"{defn}\n", {}),
        (f"   ↳ Analogy: {analogy}", {"italic": True, "color": MUTED}),
    ])

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 6. ARCHITECTURE DEEP DIVE
# ════════════════════════════════════════════════════════════════════════════
h1("6. Architecture Deep Dive")

# ── 6.1 ──
h2("6.1  The Captioning Model — ResNet18 + Transformer")

para(
    "The captioning model has two halves: an IMAGE ENCODER that turns the image into a sequence of "
    "feature vectors, and a CAPTION DECODER that generates the sentence one token at a time while looking "
    "at those vectors."
)

h3("Step 1 — Image Encoder (ResNet18)")
bullet("Input: an RGB image, resized to 224 × 224.")
bullet("ResNet18 (pretrained on ImageNet) is truncated before its global pooling layer.")
bullet("Output of the truncated backbone: a 7 × 7 × 512 feature map.")
bullet("A 1 × 1 convolution projects each spatial cell from 512-d → 256-d (our d_model).")
bullet("Flatten the 7 × 7 grid → 49 'image tokens', each a 256-d vector.")

callout(
    "Mental model",
    "The image is now a SEQUENCE of 49 vectors. The decoder will treat these the same way it treats sentence "
    "tokens — but it can only READ from them, not generate into them."
)

h3("Step 2 — Caption Decoder (Transformer)")
bullet("4-layer Transformer decoder, d_model = 256, 4 attention heads, FFN hidden size 1024.")
bullet("Token embedding: BERT WordPiece vocabulary (~30,522 tokens). [CLS] is reused as BOS, [SEP] as EOS.")
bullet("Position embedding: learned, for positions 0 to 31 (max caption length).")
bullet("Each layer has THREE sub-blocks:")
bullet("Self-attention with causal mask → each word can only attend to itself and prior words.", level=1)
bullet("Cross-attention → each word attends to the 49 image tokens (this is where the image enters the text).", level=1)
bullet("Feed-forward network → non-linear transformation.", level=1)
bullet("Output: linear projection to vocab logits, with WEIGHT TYING (shares parameters with the input embedding).")

h3("Step 3 — Generation at inference time")
code_block(
"""tokens = [BOS]                          # start with the [CLS] token
for _ in range(MAX_TOKEN_LENGTH - 1):
    logits  = decoder(tokens, image_features)
    next_id = logits[:, -1].argmax()    # greedy — pick the single best token
    if next_id == EOS:                  # stop on [SEP]
        break
    tokens.append(next_id)
caption = tokenizer.decode(tokens[1:])  # drop the BOS, decode to text""")

para(
    "This is GREEDY decoding — simple, fast, deterministic. You could replace it with BEAM SEARCH "
    "(explore the top-k candidates at each step) for slightly better captions at the cost of more compute. "
    "We chose greedy because the value-add of beam search on Flickr8k is small."
)

callout(
    "Why a Transformer here?",
    "Three reasons: (1) cross-attention is a clean way to let the decoder see image features; "
    "(2) self-attention captures long-range dependencies better than RNNs; (3) training is fully parallel — "
    "all positions in the target sequence are computed in one forward pass thanks to teacher forcing."
)

# ── 6.2 ──
h2("6.2  CLIP for Retrieval — Why It Works Without Training")

para(
    "The retrieval half uses CLIP (Contrastive Language-Image Pre-training), an off-the-shelf model from "
    "OpenAI. We TRAIN NOTHING here. Understanding why this works is the single most impressive talking point "
    "of the presentation."
)

h3("The core idea — a shared embedding space")
para(
    "CLIP has TWO encoders (one for images, one for text), each producing a 512-d vector. The training "
    "objective was: given a batch of (image, caption) pairs, push the matching image and caption embeddings "
    "TOGETHER and push everything else APART. Over 400 million pairs, this forces image and text vectors "
    "for the same concept to land in the same region of the space."
)

h3("How we use it at inference")
code_block(
"""# ONE-TIME setup (cached to disk)
image_embeddings = clip.encode_images(all_dataset_images)   # (8091, 512)
image_embeddings = normalize(image_embeddings, dim=-1)       # unit vectors

# AT QUERY TIME (milliseconds)
text_emb        = clip.encode_text(user_query)              # (1, 512)
text_emb        = normalize(text_emb, dim=-1)
similarities    = text_emb @ image_embeddings.T             # (1, 8091)
top_3_indices   = similarities.topk(3).indices""")

h3("Why cosine similarity?")
para(
    "Once vectors are normalized to unit length, the dot product equals the cosine of the angle between "
    "them. Cosine is the right choice because we care about DIRECTION (semantic meaning) not MAGNITUDE "
    "(how 'intense' a feature is)."
)

callout(
    "Important nuance",
    "After CLIP retrieves the top-3 images, OUR OWN captioning model generates a caption for each. So the "
    "result you see combines two models — CLIP for matching, our trained model for description. This is a "
    "strong demonstration of model composition."
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 7. TRAINING DETAILS
# ════════════════════════════════════════════════════════════════════════════
h1("7. Training Details & Results")

h2("Hyperparameters")
make_table(
    ["Hyperparameter", "Value", "Why"],
    [
        ["Optimizer",          "AdamW",     "Modern default; handles weight decay correctly"],
        ["Learning rate",      "3e-4",      "Karpathy's 'safe bet' for transformers"],
        ["Batch size",         "64",        "Fits on a single H100; large enough for stable gradients"],
        ["Epochs",             "20",        "Enough to see overfitting onset; lets us pick the best checkpoint"],
        ["Max token length",   "32",        "Covers >99% of Flickr8k captions"],
        ["d_model",            "256",       "Compact; matches a small dataset"],
        ["Attention heads",    "4",         "256 / 4 = 64-d per head — standard ratio"],
        ["Decoder layers",     "4",         "Enough depth without overfitting on 8k images"],
        ["FFN hidden dim",     "1024",      "4× d_model — standard Transformer ratio"],
        ["Dropout",            "0.1",       "Mild regularization"],
        ["Grad clip (max-norm)","1.0",      "Prevents loss spikes in early training"],
        ["Train / val split",  "90 / 10",   "Random, seeded for reproducibility"],
    ],
)

h2("Results")
make_table(
    ["Epoch", "Val Loss", "Val PPL", "Note"],
    [
        ["1",  "5.1747", "176.74", "Cold start"],
        ["5",  "3.3075", "27.32",  "Learning fast"],
        ["8",  "3.0613", "21.35",  "Slowing down"],
        ["10", "3.0123", "20.34",  ""],
        ["12", "2.9827", "19.74",  "★ Best — checkpoint saved"],
        ["15", "3.0186", "20.46",  "Overfitting"],
        ["20", "3.1152", "22.54",  "Overfit further"],
    ],
)

para(
    "Best validation loss: 2.9827 (perplexity 19.74) at epoch 12. Interpretation: on average the model is "
    "picking the correct next word from about 20 plausible candidates out of a 30,522-token vocabulary — "
    "much better than random.",
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 8. LOSS CURVE
# ════════════════════════════════════════════════════════════════════════════
h1("8. Loss Curve — How to Read & Explain It")

# Insert the actual figure if available
loss_img = "assets/loss_curve.png"
if os.path.exists(loss_img):
    doc.add_picture(loss_img, width=Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Validation loss across 20 epochs. Best at epoch 12 (green dot). "
                    "Red shaded region = overfitting.")
    r.font.size = Pt(9.5); r.font.italic = True; r.font.color.rgb = MUTED
else:
    para("(Run `python plot_loss.py` first to generate assets/loss_curve.png.)",
         italic=True, color=DANGER)

h2("Three phases to point at")

bullet("Phase 1 — Rapid descent (epochs 1–5). The model learns basic vocabulary and sentence structure. "
       "Loss drops from 5.17 → 3.31. This is the biggest single chunk of learning.")
bullet("Phase 2 — Slow refinement (epochs 6–12). Loss drops from 3.31 → 2.98. The model is now learning "
       "more nuanced patterns — adjectives, prepositional phrases, less common verbs.")
bullet("Phase 3 — Overfitting (epochs 13–20). Validation loss stops improving and slowly drifts upward. "
       "The model is now memorizing training-set quirks. We save the epoch-12 checkpoint and discard the rest.")

h2("One-sentence summary for the audience")
para(
    "\"The loss curve shows classic overfitting onset around epoch 12 — that's why early stopping is "
    "essential and that's the checkpoint we ship.\"",
    italic=True,
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 9. SLIDE-BY-SLIDE SPEAKER NOTES
# ════════════════════════════════════════════════════════════════════════════
h1("9. Slide-by-Slide Speaker Notes")

para(
    "For each slide: WHAT GOES ON IT (visuals + bullets) and WHAT YOU SAY (the verbatim script). "
    "Memorize one anchor sentence per slide — the rest flows from there.",
    italic=True, color=MUTED,
)

# ── Slide 1 ──
h2("Slide 1 — Title & Problem  (~30 seconds)")
h3("On the slide")
bullet("Title: 'Cross-Modal Generation — Image Captioning & Retrieval'")
bullet("Subtitle: 'ResNet18 + Transformer · CLIP · Flickr8k'")
bullet("Your name(s) and date")

h3("What you say")
para(
    "\"We built a system that bridges vision and language in two directions. Given an image, it generates a "
    "caption. Given text, it retrieves the most relevant images. A single web app, two models, one shared idea.\"",
    italic=True,
)

# ── Slide 2 ──
h2("Slide 2 — Motivation  (~45 seconds)")
h3("On the slide")
bullet("Two arrows: 'image → text' (captioning) and 'text → image' (retrieval)")
bullet("Three example applications: accessibility (screen readers), search (Google Images), content moderation")

h3("What you say")
para(
    "\"Vision and language are wildly different data types for a machine. Bridging them — cross-modal "
    "learning — enables a huge range of applications: describing images for visually impaired users, "
    "searching photo libraries by description, auto-tagging content at scale. Our project tackles both "
    "directions of this bridge.\"",
    italic=True,
)

# ── Slide 3 ──
h2("Slide 3 — Architecture Overview  (~60 seconds)")
h3("On the slide")
bullet("A two-row diagram showing the captioning pipeline (top) and retrieval pipeline (bottom)")
bullet("Annotations naming each component: ResNet18, Transformer Decoder, CLIP, cosine similarity")

h3("What you say")
para(
    "\"The system has two halves. The first — captioning — takes an image, runs it through ResNet18 to "
    "extract features, then a Transformer decoder generates a caption word by word. The second — retrieval "
    "— uses CLIP, a pretrained model that places images and text in a SHARED embedding space, and a simple "
    "cosine similarity finds the closest matches.\"",
    italic=True,
)

# ── Slide 4 ──
h2("Slide 4 — The Captioning Model  (~75 seconds, the heart of the talk)")
h3("On the slide")
bullet("Architecture diagram (the one from the README)")
bullet("Key numbers: 49 image tokens · 256-d · 4 layers · 4 heads · ~23 M parameters")

h3("What you say")
para(
    "\"Here's how the captioning model works. The image is resized to 224 by 224 and runs through ResNet18, "
    "stopped just before its final pooling layer. That gives us a 7 by 7 feature map — think of it as 49 "
    "'image tokens'. The Transformer decoder is autoregressive: it generates the caption one word at a time. "
    "At each step, it uses self-attention to look at the words it has already generated, and cross-attention "
    "to look at those 49 image tokens. Generation stops when it produces the end-of-sentence token.\"",
    italic=True,
)
para(
    "Pause for 1 second after saying 'cross-attention' — let it land. This is the keyword the technical "
    "audience will recognize.",
    italic=True, color=MUTED,
)

# ── Slide 5 ──
h2("Slide 5 — Training & Results  (~60 seconds)")
h3("On the slide")
bullet("Loss curve image (assets/loss_curve.png)")
bullet("Compact table: epoch 1, 5, 10, 12 (best), 20 with val loss & PPL")

h3("What you say")
para(
    "\"We trained on Flickr8k — 8,091 images with 5 captions each — for 20 epochs using AdamW. Teacher "
    "forcing makes training parallel: we feed the decoder the ground-truth previous words rather than its "
    "own predictions. The best validation loss came at epoch 12: 2.98, or a perplexity of about 20. From "
    "epoch 13 onward, validation loss started drifting upward — classic overfitting — so we shipped the "
    "epoch-12 checkpoint.\"",
    italic=True,
)

# ── Slide 6 ──
h2("Slide 6 — CLIP for Retrieval  (~60 seconds)")
h3("On the slide")
bullet("Diagram: text encoder + image encoder → shared 512-d space → cosine similarity")
bullet("Three bullets: 400 M training pairs · 8,091 cached embeddings · top-3 = one matrix multiply")
bullet("Optional one-line code snippet: similarity = text_emb @ image_embeddings.T")

h3("What you say")
para(
    "\"For retrieval, we TRAINED NOTHING — and that's the interesting part. OpenAI's CLIP was pretrained on "
    "400 million image-text pairs. It places images and text in the SAME 512-dimensional space, so the "
    "vector for the sentence 'a dog on the beach' lives near the vectors of actual beach-dog photos. We "
    "precomputed embeddings for all 8,091 Flickr8k images once and cached them. At query time, we encode "
    "the user's text, do a single matrix multiplication to compute cosine similarity, and pick the top "
    "three matches. Then our trained captioning model describes each of those three images — combining "
    "the two models.\"",
    italic=True,
)

# ── Slide 7 ──
h2("Slide 7 — Demo + Closing  (~90 seconds)")
h3("On the slide")
bullet("Demo URL or screenshot")
bullet("Closing bullets: three concrete future-work items")
bullet("'Thank you — questions?'")

h3("What you say (demo + close)")
para(
    "\"Let me show it briefly.\" — Run a retrieval query (e.g. 'a dog playing in the water'), point at the "
    "top-3 cards. Then drag an example image into the caption box, click Generate, read the result aloud. "
    "\"To wrap up: we trained a captioning model from scratch and combined it with a pretrained CLIP model "
    "for retrieval, all packaged in a simple Flask app. Next steps would be beam search for more fluent "
    "captions, a larger dataset like COCO for vocabulary coverage, and fine-tuning CLIP if we wanted to "
    "specialize the retrieval. Happy to take questions.\"",
    italic=True,
)

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 10. Q&A
# ════════════════════════════════════════════════════════════════════════════
h1("10. Anticipated Q&A")

para(
    "Group your answers into 'easy' (you can nail these in one sentence) and 'harder' (longer, may require "
    "thinking aloud). Prepare a fallback for questions you don't know: 'Great question — that's outside what "
    "we evaluated; my intuition is X, but I'd want to verify it before claiming it.'",
    italic=True, color=MUTED,
)

h2("Easy questions — one-sentence answers")
qa_easy = [
    ("Why a Transformer instead of an RNN?",
     "Three reasons: parallel training, better long-range dependencies, and clean cross-attention to the image features."),
    ("Why ResNet18 and not something larger?",
     "ResNet18 is light (~11 M params) and pretrained on ImageNet — it's the right size to avoid overfitting on a small "
     "8 k-image dataset."),
    ("Why didn't you fine-tune CLIP?",
     "CLIP was trained on 400 million pairs and already generalizes well. Flickr8k is tiny by comparison, so fine-tuning "
     "would risk specializing it onto noise."),
    ("What does 'teacher forcing' do?",
     "During training, the decoder sees the ground-truth previous token instead of its own prediction. This makes the "
     "whole sequence trainable in parallel and stabilizes early training."),
    ("Why early stopping?",
     "Validation loss bottomed out at epoch 12 and started rising — the model began memorizing training data. We saved "
     "epoch 12 and discarded the rest."),
    ("Why is perplexity around 20 'good'?",
     "Out of a 30 k-token vocabulary, the model is on average choosing between 20 plausible next words. That's "
     "meaningful linguistic learning, not random."),
    ("How fast is retrieval?",
     "Milliseconds. It's one matrix multiplication between the query embedding and all 8 k image embeddings."),
]
for q, a in qa_easy:
    rich([("Q: ", {"bold": True, "color": INDIGO}), (q, {"bold": True})])
    rich([("A: ", {"bold": True, "color": ACCENT}), (a, {})])

h2("Harder questions — longer answers")
qa_hard = [
    ("Greedy decoding vs. beam search — what's the trade-off?",
     "Greedy picks the single best token at each step — fast and deterministic but can get stuck. Beam search keeps "
     "the top-k candidate sequences alive at each step, exploring more of the space. Beam usually produces more fluent, "
     "less repetitive captions at maybe 5–10× the inference cost. For a demo, greedy is fine; for a paper, you'd try beam."),

    ("How would you evaluate this model rigorously?",
     "Standard captioning benchmarks: BLEU-4, CIDEr, METEOR, and SPICE. They compare the generated caption to the 5 "
     "ground-truth captions per image. We focused on loss and perplexity for this project, but adding BLEU on the val set "
     "would be the next concrete step."),

    ("What are the failure modes of your captioning model?",
     "Three common ones: (1) counting — it might say 'a dog' for two dogs; (2) rare objects — anything outside Flickr8k's "
     "everyday-scene distribution; (3) repetition — greedy decoding sometimes loops, e.g. 'a man and a man and a man'. "
     "Beam search and longer training data would help."),

    ("What does CLIP fail at?",
     "Spatial reasoning ('left of', 'on top of'), counting, and negation ('a dog that is NOT brown'). It also struggles "
     "with culturally specific or fine-grained categories. This is well documented in the CLIP paper."),

    ("Could you swap the captioning model for an LLM like GPT-4 Vision?",
     "Absolutely — modern vision-language models like GPT-4V, Gemini, or LLaVA would produce vastly better captions "
     "out of the box. The point of this project was educational: building the encoder-decoder + cross-attention "
     "pipeline from scratch teaches the mechanics. Production use would call for a pretrained VLM."),

    ("Why is the first inference request slow?",
     "Two reasons: lazy initialization of CUDA/MPS kernels and the first cache-fill of CLIP. After the first request, "
     "everything is warm and subsequent calls are an order of magnitude faster."),
]
for q, a in qa_hard:
    rich([("Q: ", {"bold": True, "color": INDIGO}), (q, {"bold": True})])
    rich([("A: ", {"bold": True, "color": ACCENT}), (a, {})])

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 11. DEMO CHECKLIST
# ════════════════════════════════════════════════════════════════════════════
h1("11. Live Demo Checklist")

h2("Before the talk")
bullet("Run `python demo.py` on the presentation laptop at least once today — confirm the checkpoint loads.")
bullet("Hit `http://127.0.0.1:7860` in the browser; verify both retrieval and captioning return real results.")
bullet("Open the page BEFORE you start presenting; don't ask the audience to wait for the first cold-cache request.")
bullet("Bookmark the page or pin the tab.")
bullet("Prepare 2 retrieval queries you've tested: one easy ('a dog on the beach'), one a bit harder ('two kids playing in the snow').")
bullet("Prepare 1 image file to drag and drop — pick one with a clean subject so the caption is impressive.")
bullet("Increase browser zoom to ~150% so the back row can see.")
bullet("Disable notifications, sleep, and screen-saver.")
bullet("Have a screenshot fallback in the slides in case the demo crashes.")

h2("During the demo (under 30 seconds)")
bullet("Type the prepared retrieval query. Point at the score and at the caption — 'CLIP found this, our model described it'.")
bullet("Drag the prepared image. Click Generate Caption. Read the output aloud.")
bullet("Move on — don't linger past 30 seconds even if the audience is enjoying it.")

h2("If something goes wrong")
bullet("Stay calm. Say 'looks like the demo's having an issue — let me show you the screenshots instead.'")
bullet("Never debug live in front of the audience. Move on, finish, return to it in Q&A if asked.")

page_break()

# ════════════════════════════════════════════════════════════════════════════
# 12. NIGHT BEFORE
# ════════════════════════════════════════════════════════════════════════════
h1("12. The Night Before — 10-Minute Review")

para("Read this list out loud the night before. If anything feels shaky, go back to the relevant section.")

h2("Can you answer these in one sentence each?")
bullet("What is cross-modal generation?")
bullet("What two halves does your system have?")
bullet("What does ResNet18 produce, and how many image tokens?")
bullet("What does cross-attention do in the decoder?")
bullet("What is teacher forcing, and why use it?")
bullet("What does perplexity mean intuitively?")
bullet("Why did training stop being useful at epoch 12?")
bullet("Why did you not need to train CLIP?")
bullet("What is cosine similarity measuring?")
bullet("What is the closing sentence of your talk?")

h2("Logistics")
bullet("Slides exported to PDF as a backup (in case PowerPoint/Keynote crashes).")
bullet("Demo verified to run on the presentation laptop.")
bullet("Loss curve image and architecture diagram included on the right slides.")
bullet("Water bottle next to the laptop.")
bullet("Phone on silent.")

callout(
    "Final pep talk",
    "You built this. You know it cold. The audience knows less than you do about this specific project — "
    "you are the expert in the room. Speak slowly, smile when you finish a section, and remember that "
    "1-second pauses feel like 5-second pauses to you but sound like 1-second pauses to them. You've got this."
)

# ════════════════════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════════════════════
out = "Cross-Modal_Presentation_Study_Guide.docx"
doc.save(out)
print(f"Saved → {out}")
print(f"Pages: ~{int(len(doc.paragraphs) / 25)} (estimate)")
